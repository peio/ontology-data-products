"""Proposed outline deck: one Graphwise-branded slide per methodology step.

Usage: python build_methodology_outline.py OUT.pptx
Needs python-pptx and the graphwise-pptx skill (graphwise_deck.py and its assets/):
set GRAPHWISE_PPTX_SKILL to the skill folder.
"""
import os
import sys

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

sys.path.insert(0, os.environ.get("GRAPHWISE_PPTX_SKILL", "."))
from graphwise_deck import (Deck, BLUE, MAGENTA, INDIGO, STEEL, INK, GREY_LT,
                            GREY_MD, WHITE, _set_bg, _text, _rect)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"

OUT = sys.argv[1]

STEPS = [
    {
        "n": 1, "name": "Domain and scope",
        "objective": "Fix what the ontology covers, what it is for, and which questions it must answer.",
        "deliverable": "Scope, out-of-scope list, six competency questions with an independent answer key",
        "did": [
            "Domain: the 50 EURO STOXX 50 constituents, with country, ICB subsector and FY2024 revenue",
            "A teaching example for workshop participants, maintained by Graphwise",
            "Six business questions (B1–B6); the answer key is computed from the CSV, never from the graph",
            "Two test tracks: is it built right (A), is it the right thing (B)",
        ],
        "rejected_label": "OUT OF SCOPE",
        "rejected": ["Country currency, membership history, FX conversion"],
        "proof": "Track B 6/6 against 45 expected rows",
        "module": "C1",
        "notes": "Ask: if the answer key were wrong, would the test notice? No — Track B is only as good as its key. Step 8 shows this happening.",
    },
    {
        "n": 2, "name": "Reuse existing ontologies",
        "objective": "Reuse well-adopted vocabularies; mint as little as possible.",
        "deliverable": "Reuse table, gap list, import plan (none, by design)",
        "did": [
            "schema.org for company, country, name and website",
            "FIBO EquityIndex for the index: schema.org has no index class",
            "ICB as a SKOS taxonomy (249 concepts), Wikidata countries, DCAT and PROV-O provenance",
            "Align, don't import: local classes sit under the reused ones, and instances carry both types",
        ],
        "rejected_label": "REJECTED",
        "rejected": ["Importing FIBO: about 132,600 triples for a 50-row dataset"],
        "proof": "A1e reuse first · A1f visible without a reasoner · A1g no redeclared terms",
        "module": "C2",
        "notes": "Three ways of reusing: alignment, direct use, reuse as data. Part 4 of the guide redoes this step with FIBO used directly.",
    },
    {
        "n": 3, "name": "Enumerate terms",
        "objective": "List every term, attribute and relation before deciding what is a class or a property.",
        "deliverable": "16 terms, their relationships, and the fate of all 9 CSV columns",
        "did": [
            "16 terms, including the ones that end up out of scope",
            "Relationships: assigned to a country, classified in a subsector, constituent of an index, reports revenue",
            "Columns: 5 mapped, 3 cross-checked and derived, 1 dropped",
        ],
        "rejected_label": "REJECTED",
        "rejected": ["Dropping a column silently: every exclusion is listed and tested"],
        "proof": "A5b every CSV column is mapped or intentionally dropped",
        "module": "C3",
        "notes": "Exercise: predict what each term becomes — class, property, taxonomy value, note, or nothing.",
    },
    {
        "n": 4, "name": "Classes and hierarchy",
        "objective": "Choose the classes and their kind-of hierarchy; declare disjointness; use SKOS for terminology.",
        "deliverable": "Three classes, each under a reused class, pairwise disjoint",
        "did": [
            "Bottom-up: Company, Country, StockIndex, from what the source rows describe",
            "Each a subclass of schema.org or FIBO, so local rules never touch terms we don't own",
            "ICB stays a SKOS taxonomy linked by skos:broader",
        ],
        "rejected_label": "REJECTED",
        "rejected": [
            "A EuroStoxx50Company class: members would migrate at every rebalancing",
            "249 OWL classes for ICB: blurs things and categories",
        ],
        "proof": "A1c no cycles · A1d every class justified · A2a no disjointness clash",
        "module": "C4",
        "notes": "Methodology §4.5: a class an instance belongs to should not change often. Membership is a relation, not a class.",
    },
    {
        "n": 5, "name": "Properties",
        "objective": "Attach each property at the most general class; declare domain, range, inverses, functional.",
        "deliverable": "Six minted and two reused properties, with naming conventions",
        "did": [
            "Link to the 8-digit subsector only; industry, sector and supersector are derived",
            "One country property: the index-assigned country (Stellantis: domiciled NL, assigned IT)",
            "Revenue next to its reporting currency; an inverse pair for index membership",
        ],
        "rejected_label": "REJECTED",
        "rejected": [
            "Storing every ICB level: redundant, and can contradict itself",
            "Separate domicile and headquarters: no question needs them",
        ],
        "proof": "A1a labels and definitions · A1b domain and range · A5c roll-up consistent",
        "module": "C5",
        "notes": "Naming: UpperCamelCase classes, lowerCamelCase properties. Exception to discuss: revenueFY2024 puts the year in the name.",
    },
    {
        "n": 6, "name": "Facets and constraints",
        "objective": "Cardinalities, value types, ranges and patterns as SHACL; OWL only where reasoning needs it.",
        "deliverable": "SHACL shapes, each with a plain-language message",
        "did": [
            "OWL: disjoint classes, functional properties, the inverse pair",
            "SHACL: exactly one of each fact, revenue ≥ 0, a three-letter currency, a real subsector",
            "Three planted broken companies produce exactly three violations",
        ],
        "rejected_label": "KEY INSIGHT",
        "rejected": ["Give a company two countries: OWL silently merges them; only SHACL raises an error"],
        "proof": "A3a clean data conforms · A3b three planted violations caught",
        "module": "C6",
        "notes": "OWL works in an open world: a missing fact is unknown, not false, so it cannot say 'must have'. SHACL checks the data as it is.",
    },
    {
        "n": 7, "name": "Instances and testing",
        "objective": "Populate the ontology with data and test it against the competency questions.",
        "deliverable": "50 companies, 650 triples, a test report per question",
        "did": [
            "Identifiers gwi:company-<slug>; Wikidata IRIs for countries",
            "Provisional ICB subsectors; the mapping refuses any that don't roll up to the stated industry",
            "Every company traces back to both source files",
        ],
        "rejected_label": "RESULT",
        "rejected": ["Track A 22/22 · Track B 6/6"],
        "proof": "A5a 50 companies · A5c 50 of 50 consistent · A6a provenance",
        "module": "C7",
        "notes": "Exercise: give SAP a Financials subsector and watch the mapping refuse to run.",
    },
    {
        "n": 8, "name": "Refine and iterate",
        "objective": "Fix what testing and review find, returning to the step where each problem began.",
        "deliverable": "Two revisions; version 2.0.0",
        "did": [
            "Revision 1, back to Steps 2 and 7: no schema.org types without a reasoner. Fix: dual typing, two new tests",
            "Revision 2, back to Steps 1, 3 and 5: EUR and USD were summed. Fix: a reporting-currency property",
            "Also fixed: legal names that were short names, a namespace shared with other ontologies",
        ],
        "rejected_label": "LESSON",
        "rejected": ["Every test passed both times. A green run proves only what the tests ask"],
        "proof": "Track A 20/20 → 22/22 · ontology 1.0.0 → 2.0.0",
        "module": "C8a, C8b",
        "notes": "TotalEnergies (195.6, in USD) was the 'largest company' — and the test passed, because the answer key made the same mistake.",
    },
    {
        "n": 9, "name": "Document and deploy",
        "objective": "Make the result usable without this conversation: documentation, versioning, deployment, maintenance.",
        "deliverable": "Reference docs, semantic-layer explorer, GraphDB load notes",
        "did": [
            "Version 2.0.0 with prior version, licence CC BY 4.0",
            "An explorer page for people who will never read Turtle or SPARQL",
            "GraphDB with OWL2-RL; shapes as a SHACL graph",
            "Maintenance: full pipeline on every change; an assumption register",
        ],
        "rejected_label": "TRAVELS WITH THE DATA",
        "rejected": ["Provenance and caveats: figures approximate, subsectors provisional"],
        "proof": "A6a–A6c provenance, versions, caveats · B6 where a figure comes from",
        "module": "C9",
        "notes": "Reproducible: with pinned dependencies every generated file rebuilds byte for byte.",
    },
]


def dot(slide, x, y, d, color):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def label(slide, x, y, w, text, color=STEEL):
    _text(slide, x, y, w, 0.3, [[(text, {"size": 11, "bold": True, "color": color, "font": FONT})]])


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


d = Deck()
FONT = d.font

s = d.cover("Building a tested ontology, step by step",
            subtitle="The nine steps of the Graphwise ontology-engineer method, on the EURO STOXX 50",
            presenter="Peio Popov", date="AI Week Frankfurt · October 2026")
notes(s, "Proposed outline: one slide per methodology step. Detail for each slide is in tutorial/guide.")

# Overview: the nine steps as a process row.
s = d._slide()
_set_bg(s, WHITE)
d._title(s, "Nine steps, one tested ontology")
names = ["Domain and scope", "Reuse", "Enumerate terms", "Classes", "Properties",
         "Facets and constraints", "Instances and test", "Refine and iterate", "Document and deploy"]
x0, step_w, cy, cd = 0.9, 11.53 / 9, 2.35, 0.78
line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0 + step_w / 2), Inches(cy + cd / 2),
                              Inches(x0 + step_w * 8.5), Inches(cy + cd / 2))
line.line.color.rgb = GREY_MD
line.line.width = Pt(1.5)
for i, name in enumerate(names):
    cx = x0 + step_w * i + (step_w - cd) / 2
    c = dot(s, cx, cy, cd, BLUE)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(i + 1)
    r.font.size, r.font.bold, r.font.name = Pt(22), True, FONT
    r.font.color.rgb = WHITE
    _text(s, x0 + step_w * i + 0.05, cy + cd + 0.25, step_w - 0.1, 0.9,
          [[(name, {"size": 13, "bold": True, "color": INDIGO, "font": FONT})]],
          align=PP_ALIGN.CENTER, line_spacing=1.05)
_text(s, 0.9, 4.75, 11.5, 1.2,
      [[("Steps 4 and 5 are interleaved. Steps 7 and 8 loop until the tests and the review both pass. "
         "Every decision, and what was rejected, goes in the decision log.",
         {"size": 15, "color": STEEL, "font": FONT})]], line_spacing=1.1)
d._footer(s)
notes(s, "Source: skills/graphwise-ontology-engineer/SKILL.md (Workflow), based on Noy and McGuinness, Ontology Development 101.")

for st in STEPS:
    accent = BLUE if st["n"] % 2 else MAGENTA
    s = d._slide()
    _set_bg(s, WHITE)
    d._title(s, f"Step {st['n']} · {st['name']}")

    # Left panel: number, objective, deliverable.
    _rect(s, 0.9, 1.75, 3.9, 4.85, GREY_LT)
    _text(s, 1.25, 1.95, 3.2, 1.0,
          [[(f"{st['n']:02d}", {"size": 54, "bold": True, "color": accent, "font": FONT})]])
    label(s, 1.25, 3.05, 3.2, "OBJECTIVE")
    _text(s, 1.25, 3.35, 3.2, 1.6,
          [[(st["objective"], {"size": 15, "color": INK, "font": FONT})]], line_spacing=1.08)
    label(s, 1.25, 5.0, 3.2, "DELIVERABLE")
    _text(s, 1.25, 5.3, 3.2, 1.2,
          [[(st["deliverable"], {"size": 13, "color": INK, "font": FONT})]], line_spacing=1.08)

    # Right: what we did, then rejected / insight, in one flowing text frame; then proof.
    rx, rw = 5.3, 7.1
    box = s.shapes.add_textbox(Inches(rx), Inches(1.8), Inches(rw), Inches(4.0))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = [True]

    def para(text, *, size, bold=False, bullet=None, before=0, after=6):
        p = tf.paragraphs[0] if first[0] else tf.add_paragraph()
        first[0] = False
        p.line_spacing = 1.04
        p.space_before, p.space_after = Pt(before), Pt(after)
        r = p.add_run()
        r.text = text
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = INDIGO if bold else INK
        if bullet is not None:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.3)))
            pPr.set("indent", str(-Inches(0.3)))
            hexcol = str(bullet)
            pPr.append(etree.fromstring(
                f'<a:buClr xmlns:a="{A}"><a:srgbClr val="{hexcol}"/></a:buClr>'))
            pPr.append(etree.fromstring(f'<a:buSzPct xmlns:a="{A}" val="90000"/>'))
            pPr.append(etree.fromstring(f'<a:buFont xmlns:a="{A}" typeface="Arial"/>'))
            pPr.append(etree.fromstring(f'<a:buChar xmlns:a="{A}" char="●"/>'))

    para("WHAT WE DID", size=11, bold=True, after=8)
    for item in st["did"]:
        para(item, size=15, bullet=accent)
    para(st["rejected_label"], size=11, bold=True, before=10, after=8)
    for item in st["rejected"]:
        para(item, size=15, bullet=GREY_MD)
    label(s, rx, 5.95, rw, "PROOF")
    _text(s, rx, 6.25, rw, 0.35,
          [[(st["proof"], {"size": 12, "color": STEEL, "font": FONT})]])
    label(s, 1.25, 6.15, 3.2, f"GUIDE  {st['module']}")
    d._footer(s)
    notes(s, f"Guide module {st['module']}. {st['notes']}")

s = d.two_column("Back to Step 2: the same data with FIBO",
                 "Light reuse · ontology 2.0.0",
                 ["Local classes under schema.org and FIBO",
                  "Revenue: a number plus a currency",
                  "Wikidata countries",
                  "9 minted terms · Track A 22/22, B 6/6"],
                 "Deep reuse · FIBO variant",
                 ["28 FIBO, Commons and LCC terms, used as is",
                  "Revenue node: currencies can't mix",
                  "LCC countries; a flat index link",
                  "5 minted terms · Track A 26/26, B 6/6"])
notes(s, "Guide Part 4 (F0–F7). Neither model is right: light reuse is quicker and easier to read; deep reuse speaks the industry's language, at the price of inherited expectations and a billions caveat.")

s = d.closing("Every decision recorded. Every answer tested.",
              contacts={"Web": "graphwise.ai", "Code": "github.com/peio/ontology-data-products"})
d.save(OUT)

print("saved", OUT)
